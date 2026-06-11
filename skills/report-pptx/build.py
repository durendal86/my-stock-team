"""report-pptx 스킬 렌더러.

reports/{종목명}.md 를 읽어 디자인된 reports/{종목명}.pptx 를 생성한다.

슬라이드 순서(고정): 표지 → 종목 개요 → 재무 요약 → 가격/추세 → 뉴스·심리 → 리스크 → 한 줄 종합
디자인: 증권사 리서치 톤 — 다크 차콜 표지, 섹션 키커(01|OVERVIEW), KB 옐로우(#FFBC00)는
포인트로만 절제 사용, 차콜 헤더 표(숫자 우측정렬), 헤어라인 푸터 + 페이지 번호. 맑은 고딕 고정.

입력 .md 포맷:
    ---
    종목명: 삼성전자
    종목코드: 005930        # (선택) 있으면 FinanceDataReader로 가격 차트 자동 생성
    작성일: 2026-06-11
    ---

    ## 종목 개요
    ...본문...

    ## 재무 요약
    | 항목 | 2023 | 2024 | 2025 |
    |---|---|---|---|
    | 매출액 | ... | ... | ... |
    (출처: DART, 2025 사업보고서 / 기준일 2025-12-31)

    ## 가격/추세
    ![chart](상대경로.png)   # (선택) 직접 이미지. 없고 종목코드가 있으면 자동 차트.
    ...본문...

    ## 뉴스·심리
    - ...

    ## 리스크
    - ...

    ## 한 줄 종합
    한 문장 종합 의견 (매매 단정 금지)

사용: python build.py reports/삼성전자.md
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

import yaml
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

# --- 디자인 토큰 -----------------------------------------------------------
KB_YELLOW = RGBColor(0xFF, 0xBC, 0x00)
CHARCOAL = RGBColor(0x1C, 0x1E, 0x23)   # 표지·표 헤더 배경
INK = RGBColor(0x26, 0x28, 0x2D)        # 본문 텍스트
GRAY = RGBColor(0x6B, 0x6E, 0x76)       # 보조 텍스트
FAINT = RGBColor(0xA9, 0xAC, 0xB3)      # 푸터·캡션
HAIR = RGBColor(0xDD, 0xDE, 0xE1)       # 헤어라인
PANEL = RGBColor(0xF6, 0xF6, 0xF7)      # 패널 배경
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "맑은 고딕"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
MARGIN = Inches(0.85)
CONTENT_W = SLIDE_W - 2 * MARGIN

BRAND = "STOCK TEAM  RESEARCH"

# 매매 단정 금지 가드레일 — 발견 시 경고(렌더는 진행하되 운용역에 알림)
BANNED = ["매수 추천", "매도 추천", "사야 합니다", "팔아야 합니다", "목표주가", "목표가",
          "강력 매수", "비중 확대 추천", "비중 축소 추천", "buy", "sell rating"]

# 섹션 별칭 → 표준 키
SECTION_ALIASES = {
    "종목 개요": "overview", "개요": "overview", "기업 개요": "overview",
    "재무 요약": "financials", "재무": "financials", "재무 요약(표)": "financials",
    "가격/추세": "price", "가격": "price", "추세": "price", "차트": "price", "가격·추세": "price",
    "뉴스·심리": "news", "뉴스": "news", "심리": "news", "뉴스/심리": "news", "뉴스 심리": "news",
    "리스크": "risk", "리스크 요인": "risk", "위험": "risk",
    "한 줄 종합": "summary", "한줄 종합": "summary", "종합": "summary",
    "종합의견": "summary", "한 줄 요약": "summary",
}

# 콘텐츠 슬라이드 키커 (번호 | 영문 라벨)
KICKERS = {
    "overview": ("01", "COMPANY OVERVIEW"),
    "financials": ("02", "FINANCIAL SUMMARY"),
    "price": ("03", "PRICE & TREND"),
    "news": ("04", "NEWS & SENTIMENT"),
    "risk": ("05", "RISK FACTORS"),
    "summary": ("06", "BOTTOM LINE"),
}


# --- 폰트 헬퍼 (latin/ea/cs 모두 맑은 고딕 → 한글 깨짐 방지) -----------------
def _apply_font(run, size, bold, color, spacing=None):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = FONT
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", FONT)
    if spacing is not None:                      # 자간 (1/100pt)
        rPr.set("spc", str(spacing))


def _para(tf, text, size=14, bold=False, color=INK, level=0, align=PP_ALIGN.LEFT,
          first=False, space_after=6, spacing=None, line=None):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.level = level
    p.alignment = align
    p.space_after = Pt(space_after)
    if line is not None:
        p.line_spacing = line
    run = p.add_run()
    run.text = text
    _apply_font(run, size, bold, color, spacing=spacing)
    return p


def _textbox(slide, left, top, width, height):
    box = slide.shapes.add_textbox(left, top, width, height)
    box.text_frame.word_wrap = True
    return box.text_frame


def _blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def _rect(slide, left, top, width, height, color):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


# --- 콘텐츠 슬라이드 공통 프레임 --------------------------------------------
def _header(slide, key, title, name):
    """키커(번호|라벨) + 타이틀 + 타이틀 하단 헤어라인."""
    no, label = KICKERS[key]
    tf = _textbox(slide, MARGIN, Inches(0.5), CONTENT_W, Inches(0.32))
    p = tf.paragraphs[0]
    r1 = p.add_run(); r1.text = no
    _apply_font(r1, 11, True, KB_YELLOW, spacing=200)
    r2 = p.add_run(); r2.text = f"   |   {label}"
    _apply_font(r2, 11, True, GRAY, spacing=200)

    tf = _textbox(slide, MARGIN, Inches(0.85), CONTENT_W - Inches(2.4), Inches(0.75))
    _para(tf, title, size=27, bold=True, color=INK, first=True)

    # 우측 상단 종목명 (작게)
    tf = _textbox(slide, SLIDE_W - MARGIN - Inches(2.4), Inches(0.95), Inches(2.4), Inches(0.4))
    _para(tf, name, size=11, bold=True, color=FAINT, align=PP_ALIGN.RIGHT, first=True)

    _rect(slide, MARGIN, Inches(1.62), CONTENT_W, Pt(1.4), HAIR)
    _rect(slide, MARGIN, Inches(1.62), Inches(0.55), Pt(2.6), KB_YELLOW)


def _footer(slide, asof, page=None):
    _rect(slide, MARGIN, SLIDE_H - Inches(0.52), CONTENT_W, Pt(1), HAIR)
    tf = _textbox(slide, MARGIN, SLIDE_H - Inches(0.45), CONTENT_W - Inches(0.8), Inches(0.32))
    _para(tf, f"{BRAND}   ·   학습용 분석 자료   ·   모든 수치는 표기된 출처·기준일 기준   ·   기준일 {asof}",
          size=8.5, color=FAINT, first=True, spacing=60)
    if page is not None:
        tf = _textbox(slide, SLIDE_W - MARGIN - Inches(0.7), SLIDE_H - Inches(0.45),
                      Inches(0.7), Inches(0.32))
        _para(tf, f"{page:02d}", size=9, bold=True, color=GRAY, align=PP_ALIGN.RIGHT, first=True)


# --- 마크다운 파싱 ----------------------------------------------------------
def parse_md(text):
    meta = {}
    body = text
    m = re.match(r"^\s*---\n(.*?)\n---\n", text, re.S)
    if m:
        meta = yaml.safe_load(m.group(1)) or {}
        body = text[m.end():]

    sections = {}
    cur, buf = None, []
    for line in body.splitlines():
        h = re.match(r"^##\s+(.*\S)\s*$", line)
        if h:
            if cur:
                sections[cur] = "\n".join(buf).strip()
            raw = h.group(1).strip()
            cur = SECTION_ALIASES.get(raw, raw)
            buf = []
        elif cur:
            buf.append(line)
    if cur:
        sections[cur] = "\n".join(buf).strip()
    return meta, sections


def parse_table(block):
    """첫 마크다운 표를 (헤더, 행들)로 반환. 없으면 None."""
    rows = []
    for line in block.splitlines():
        if line.strip().startswith("|"):
            cells = [c.strip() for c in line.strip().strip("|").split("|")]
            if set("".join(cells)) <= set("-: "):   # 구분선 행
                continue
            rows.append(cells)
    if not rows:
        return None
    return rows[0], rows[1:]


def find_image(block, base_dir):
    m = re.search(r"!\[[^\]]*\]\(([^)]+)\)", block)
    if not m:
        return None
    p = (base_dir / m.group(1)).resolve()
    return p if p.exists() else None


def strip_md(line):
    line = re.sub(r"!\[[^\]]*\]\([^)]+\)", "", line)      # 이미지 제거
    line = re.sub(r"\*\*(.+?)\*\*", r"\1", line)            # 볼드
    line = re.sub(r"`(.+?)`", r"\1", line)
    return line.strip()


def text_lines(block):
    """표/이미지 줄을 빼고 본문 줄 목록(불릿 레벨 포함) 반환."""
    out = []
    for line in block.splitlines():
        s = line.strip()
        if not s or s.startswith("|") or s.startswith("!["):
            continue
        bullet = bool(re.match(r"^[-*]\s+", s))
        s = re.sub(r"^[-*]\s+", "", s)
        s = re.sub(r"^\d+\.\s+", "", s)
        out.append((strip_md(s), 1 if bullet else 0))
    return out


def _body_block(slide, lines, top, size=13.5, width=None, gap=10):
    """본문 영역 — 불릿은 옐로우 틱 마커로."""
    width = width or CONTENT_W
    tf = _textbox(slide, MARGIN, top, width, SLIDE_H - top - Inches(0.7))
    for i, (txt, lvl) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        p.line_spacing = 1.15
        if lvl:
            mark = p.add_run(); mark.text = "—  "
            _apply_font(mark, size, True, KB_YELLOW)
        run = p.add_run(); run.text = txt
        _apply_font(run, size, False, INK)
    return tf


# --- 차트 생성 (선택) -------------------------------------------------------
def make_chart(code, name, out_dir, asof):
    try:
        import FinanceDataReader as fdr
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
        # matplotlib는 맑은 고딕을 영문 패밀리명 'Malgun Gothic'으로 인식한다
        # (PPTX 텍스트용 FONT='맑은 고딕'과 다름). 차트 내 한글 깨짐 방지.
        plt.rcParams["font.family"] = "Malgun Gothic"
        plt.rcParams["axes.unicode_minus"] = False

        df = fdr.DataReader(str(code), "2025-09-01", asof)
        if df.empty:
            return None, None
        df["MA20"] = df["Close"].rolling(20).mean()
        df["MA60"] = df["Close"].rolling(60).mean()
        view = df.iloc[-126:]

        fig, ax = plt.subplots(figsize=(8.6, 4.1))
        fig.patch.set_facecolor("white")
        ax.set_facecolor("white")
        ax.plot(view.index, view["Close"], label="종가", color="#26282D", linewidth=1.7)
        ax.plot(view.index, view["MA20"], label="MA20", color="#FFBC00", linewidth=1.4)
        ax.plot(view.index, view["MA60"], label="MA60", color="#A9ACB3", linewidth=1.2)
        for side in ("top", "right"):
            ax.spines[side].set_visible(False)
        for side in ("left", "bottom"):
            ax.spines[side].set_color("#DDDEE1")
        ax.set_title(f"{name} ({code})  ·  최근 6개월 종가/이동평균", fontsize=11,
                     loc="left", color="#26282D", pad=12)
        ax.legend(loc="upper left", fontsize=9, frameon=False)
        ax.grid(axis="y", alpha=0.35, color="#DDDEE1", linewidth=0.7)
        ax.grid(axis="x", visible=False)
        ax.tick_params(labelsize=8, colors="#6B6E76", length=0)
        fig.autofmt_xdate()
        fig.tight_layout()
        cdir = out_dir / ".charts"
        cdir.mkdir(parents=True, exist_ok=True)
        path = cdir / f"{code}.png"
        fig.savefig(path, dpi=150)
        plt.close(fig)
        last = view.index[-1].strftime("%Y-%m-%d")
        return path, last
    except Exception as e:  # 네트워크/라이브러리 문제 시 차트 생략
        print(f"[report-pptx] 차트 생성 건너뜀: {e}")
        return None, None


# --- 슬라이드 빌더 ----------------------------------------------------------
def slide_cover(prs, name, meta, asof):
    s = _blank(prs)
    _rect(s, 0, 0, SLIDE_W, SLIDE_H, CHARCOAL)
    _rect(s, 0, 0, Inches(0.16), SLIDE_H, KB_YELLOW)         # 좌측 옐로우 스파인

    tf = _textbox(s, MARGIN, Inches(0.75), CONTENT_W, Inches(0.4))
    _para(tf, BRAND, size=12, bold=True, color=KB_YELLOW, first=True, spacing=420)

    tf = _textbox(s, MARGIN, Inches(2.55), CONTENT_W, Inches(2.4))
    _para(tf, "Company Report", size=15, bold=False, color=FAINT, first=True,
          spacing=240, space_after=14)
    code = str(meta.get("종목코드", "")).strip()
    title = f"{name}" + (f"  ({code})" if code else "")
    _para(tf, title, size=44, bold=True, color=WHITE, space_after=10)

    _rect(s, MARGIN, Inches(4.45), Inches(2.6), Pt(2.2), KB_YELLOW)

    tf = _textbox(s, MARGIN, Inches(4.75), CONTENT_W, Inches(1.0))
    _para(tf, f"작성일  {asof}", size=13, color=WHITE, first=True, space_after=4)
    _para(tf, "펀더멘털 · 가격/추세 · 뉴스·심리 · 리스크 종합", size=11.5, color=FAINT)

    tf = _textbox(s, MARGIN, SLIDE_H - Inches(0.85), CONTENT_W, Inches(0.5))
    _para(tf, "본 자료는 학습용 분석이며, 매수·매도 등 투자 판단의 근거로 사용될 수 없습니다.",
          size=9, color=FAINT, first=True)


def slide_text(prs, key, title, block, name, asof, page):
    s = _blank(prs)
    _header(s, key, title, name)
    lines = text_lines(block) or [("(내용 없음)", 0)]
    _body_block(s, lines, top=Inches(2.0), size=13.5)
    _footer(s, asof, page)
    return s


def slide_financials(prs, block, name, asof, page):
    s = _blank(prs)
    _header(s, "financials", "재무 요약", name)
    parsed = parse_table(block)
    top = Inches(2.05)
    if parsed:
        header, rows = parsed
        # 최근 3개년 우선: 첫 열(항목) + 마지막 3개 데이터 열
        if len(header) > 4:
            keep = [0] + list(range(len(header) - 3, len(header)))
            header = [header[i] for i in keep]
            rows = [[r[i] for i in keep if i < len(r)] for r in rows]
        # 오버플로 방지: 표 높이 한계 내로 행 수/폰트 조정
        max_h = SLIDE_H - top - Inches(1.35)
        n = len(rows) + 1
        fsize, truncated = 13, False
        if n > 7:
            fsize = 11.5
        if n > 11:
            fsize = 10
        row_h = Inches(0.5 if n <= 7 else 0.42)
        max_rows = int(max_h / row_h)
        if n > max_rows:
            rows = rows[: max_rows - 1]
            truncated = True
            n = len(rows) + 1
        _draw_table(s, header, rows, top, fsize, row_h)
        note_top = top + row_h * n + Inches(0.22)
        notes = [ln for ln in block.splitlines() if "출처" in ln or "기준일" in ln]
        tf = _textbox(s, MARGIN, note_top, CONTENT_W, Inches(1.2))
        first = True
        # 표 아래 본문 불릿(있으면 1~2줄)
        body = [t for t, lvl in text_lines(block)][:2]
        for b in body:
            _para(tf, b, size=11.5, color=INK, first=first, space_after=4)
            first = False
        if notes:
            _para(tf, strip_md(notes[0]).strip("()"), size=9.5, color=FAINT, first=first)
            first = False
        if truncated:
            _para(tf, "※ 행이 많아 일부만 표기했습니다(원문 .md 참조).", size=9.5,
                  color=FAINT, first=first)
    else:
        _body_block(s, text_lines(block) or [("(표 없음)", 0)], top=top)
    _footer(s, asof, page)


def _draw_table(slide, header, rows, top, fsize, row_h):
    n_rows, n_cols = len(rows) + 1, len(header)
    width = CONTENT_W
    tbl_shape = slide.shapes.add_table(n_rows, n_cols, MARGIN, top, width, row_h * n_rows)
    tbl = tbl_shape.table
    # 기본 스타일 밴딩 제거
    tbl.first_row = False
    tbl.horz_banding = False
    # 첫 열 넓게
    tbl.columns[0].width = Emu(int(width * 0.34))
    rest = int((width - tbl.columns[0].width) / max(1, n_cols - 1))
    for c in range(1, n_cols):
        tbl.columns[c].width = Emu(rest)
    for c, h in enumerate(header):
        _fill_cell(tbl.cell(0, c), h, fsize, bold=True, bg=CHARCOAL,
                   color=WHITE, align=PP_ALIGN.LEFT if c == 0 else PP_ALIGN.RIGHT)
    for r, row in enumerate(rows, start=1):
        bg = WHITE if r % 2 else PANEL
        for c in range(n_cols):
            val = row[c] if c < len(row) else ""
            _fill_cell(tbl.cell(r, c), val, fsize, bold=(c == 0), bg=bg, color=INK,
                       align=PP_ALIGN.LEFT if c == 0 else PP_ALIGN.RIGHT)
    # 헤더 아래 옐로우 룰
    _rect(slide, MARGIN, top + row_h, width, Pt(2), KB_YELLOW)


def _fill_cell(cell, text, size, bold, bg, color, align=PP_ALIGN.LEFT):
    cell.fill.solid()
    cell.fill.fore_color.rgb = bg
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.margin_left = Pt(10)
    cell.margin_right = Pt(10)
    cell.margin_top = Pt(2)
    cell.margin_bottom = Pt(2)
    tf = cell.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    _apply_font(run, size, bold, color)


def slide_price(prs, block, base_dir, meta, name, asof, page):
    s = _blank(prs)
    _header(s, "price", "가격 / 추세", name)
    img = find_image(block, base_dir)
    chart_asof = None
    if img is None and meta.get("종목코드"):
        img, chart_asof = make_chart(meta["종목코드"], meta.get("종목명", ""), base_dir, asof)
    lines = text_lines(block)
    if img:
        s.shapes.add_picture(str(img), MARGIN, Inches(2.0), width=Inches(7.5))
        # 우측 코멘트 패널
        px = MARGIN + Inches(7.8)
        pw = SLIDE_W - px - MARGIN
        _rect(s, px, Inches(2.0), pw, Inches(4.4), PANEL)
        _rect(s, px, Inches(2.0), pw, Pt(2.2), KB_YELLOW)
        tf = _textbox(s, px + Inches(0.22), Inches(2.2), pw - Inches(0.44), Inches(4.0))
        _para(tf, "TREND NOTE", size=10, bold=True, color=GRAY, first=True,
              spacing=240, space_after=10)
        for txt, lvl in lines:
            p = tf.add_paragraph()
            p.space_after = Pt(8)
            p.line_spacing = 1.18
            run = p.add_run(); run.text = txt
            _apply_font(run, 11, False, INK)
        if chart_asof:
            p = tf.add_paragraph()
            run = p.add_run(); run.text = f"출처: FinanceDataReader · 기준일 {chart_asof}"
            _apply_font(run, 8.5, False, FAINT)
    else:
        _body_block(s, lines or [("(차트/내용 없음)", 0)], top=Inches(2.0))
    _footer(s, asof, page)


def slide_summary(prs, block, name, asof, page):
    s = _blank(prs)
    _header(s, "summary", "한 줄 종합", name)
    line = next((strip_md(l) for l in block.splitlines()
                 if strip_md(l) and not l.strip().startswith(("|", "!["))), "(종합 없음)")
    # 인용 블록: 좌측 옐로우 스파인 + 패널
    _rect(s, MARGIN, Inches(2.7), CONTENT_W, Inches(2.3), PANEL)
    _rect(s, MARGIN, Inches(2.7), Inches(0.12), Inches(2.3), KB_YELLOW)
    tf = _textbox(s, MARGIN + Inches(0.55), Inches(2.95), CONTENT_W - Inches(1.1), Inches(1.8))
    _para(tf, line, size=20, bold=True, color=INK, first=True, line=1.3, space_after=14)
    _para(tf, "※ 판단 근거 요약입니다. 매수·매도·목표가는 단정하지 않으며, 최종 판단은 운용역이 합니다.",
          size=11, color=GRAY)
    _footer(s, asof, page)


# --- 엔트리포인트 -----------------------------------------------------------
def build(md_path: Path) -> Path:
    text = md_path.read_text(encoding="utf-8")
    meta, sec = parse_md(text)
    name = meta.get("종목명") or md_path.stem
    asof = str(meta.get("작성일", "")).strip() or "기준일 미기재"

    for key, block in sec.items():
        for bad in BANNED:
            if bad.lower() in block.lower():
                print(f"[report-pptx] ⚠ 가드레일 경고: '{bad}' 표현이 '{key}' 섹션에 있습니다. "
                      "매매 단정 표현은 제거를 권장합니다.")

    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    slide_cover(prs, name, meta, asof)
    slide_text(prs, "overview", "종목 개요", sec.get("overview", ""), name, asof, 2)
    slide_financials(prs, sec.get("financials", ""), name, asof, 3)
    slide_price(prs, sec.get("price", ""), md_path.parent, meta, name, asof, 4)
    slide_text(prs, "news", "뉴스 · 심리", sec.get("news", ""), name, asof, 5)
    slide_text(prs, "risk", "리스크", sec.get("risk", ""), name, asof, 6)
    slide_summary(prs, sec.get("summary", ""), name, asof, 7)

    out = md_path.with_suffix(".pptx")
    prs.save(out)
    return out


def main():
    args = [a for a in sys.argv[1:] if a != "--pdf"]
    to_pdf = "--pdf" in sys.argv[1:]
    if not args:
        print("사용법: python build.py reports/{종목명}.md [--pdf]")
        sys.exit(1)
    md_path = Path(args[0]).resolve()
    if not md_path.exists():
        print(f"입력 파일이 없습니다: {md_path}")
        sys.exit(1)
    out = build(md_path)
    print(f"리포트 생성 완료: {out}")

    if to_pdf:  # pdf-convert 스킬로 위임
        import importlib.util
        conv_path = Path(__file__).resolve().parent.parent / "pdf-convert" / "convert.py"
        spec = importlib.util.spec_from_file_location("pdf_convert", conv_path)
        pc = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(pc)
        sys.exit(pc.convert([out], None))


if __name__ == "__main__":
    main()
